from __future__ import annotations

import os, re, math, textwrap
from pathlib import Path
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent
OUT = ROOT / "output"
OUT.mkdir(exist_ok=True)
SOURCE = ROOT / "r15_source.txt"
FIG = OUT / "Figure_Overall_Strategy.png"
DOCX = OUT / "R15_Research_Strategy_Final_With_Strategy_Figure.docx"

# ---------- Font handling ----------
FONT_CANDIDATES = [
    "/usr/share/fonts/truetype/msttcorefonts/Arial.ttf",
    "/usr/share/fonts/truetype/msttcorefonts/arial.ttf",
    "/usr/share/fonts/truetype/msttcorefonts/Arial_Bold.ttf",
    "/usr/share/fonts/truetype/msttcorefonts/arialbd.ttf",
    "/usr/share/fonts/truetype/liberation2/LiberationSans-Regular.ttf",
    "/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf",
]
BOLD_CANDIDATES = [
    "/usr/share/fonts/truetype/msttcorefonts/Arial_Bold.ttf",
    "/usr/share/fonts/truetype/msttcorefonts/arialbd.ttf",
    "/usr/share/fonts/truetype/liberation2/LiberationSans-Bold.ttf",
    "/usr/share/fonts/truetype/liberation/LiberationSans-Bold.ttf",
]
def first_existing(cands):
    for p in cands:
        if os.path.exists(p):
            return p
    return cands[-1]
FONT_REG = first_existing(FONT_CANDIDATES)
FONT_BOLD = first_existing(BOLD_CANDIDATES)

def fnt(size, bold=False):
    return ImageFont.truetype(FONT_BOLD if bold else FONT_REG, size=size)

# ---------- Figure ----------
W, H = 2400, 1350
img = Image.new("RGB", (W, H), "white")
d = ImageDraw.Draw(img)
BLACK=(25,25,25); GRAY=(130,130,130); LIGHT=(245,245,245); MID=(220,220,220)
BLUE=(38,130,255); GREEN=(38,180,90); PURPLE=(135,70,225); ORANGE=(245,155,35)
RED=(235,48,50); TEAL=(40,160,155); SKY=(60,165,240)
LW=2

def box(x0,y0,x1,y1,fill="white",outline=BLACK,width=LW,r=0):
    if r:
        d.rounded_rectangle([x0,y0,x1,y1],radius=r,fill=fill,outline=outline,width=width)
    else:
        d.rectangle([x0,y0,x1,y1],fill=fill,outline=outline,width=width)

def center_text(text, xy, font, fill=BLACK, anchor="mm"):
    d.text(xy,text,font=font,fill=fill,anchor=anchor)

def text_in(text, x0,y0,x1,y1,font,fill=BLACK,align="left",pad=12,spacing=4):
    maxw=x1-x0-2*pad
    words=text.split()
    lines=[]; cur=""
    for w in words:
        test=(cur+" "+w).strip()
        if d.textbbox((0,0),test,font=font)[2] <= maxw:
            cur=test
        else:
            if cur: lines.append(cur)
            cur=w
    if cur: lines.append(cur)
    lineh=font.size+spacing
    total=len(lines)*lineh-spacing
    yy=y0+(y1-y0-total)/2
    for line in lines:
        if align=="center":
            xx=(x0+x1)/2; anch="ma"
        elif align=="right":
            xx=x1-pad; anch="ra"
        else:
            xx=x0+pad; anch="la"
        d.text((xx,yy),line,font=font,fill=fill,anchor=anch)
        yy+=lineh

def arrow(x1,y1,x2,y2,color=BLACK,width=3,head=10):
    d.line((x1,y1,x2,y2),fill=color,width=width)
    ang=math.atan2(y2-y1,x2-x1)
    p1=(x2-head*math.cos(ang-math.pi/6),y2-head*math.sin(ang-math.pi/6))
    p2=(x2-head*math.cos(ang+math.pi/6),y2-head*math.sin(ang+math.pi/6))
    d.polygon([p1,p2,(x2,y2)],fill=color)

def seq_boxes(x,y,scale=1.0):
    vals=[4,1,3,2,4]; cols={4:BLUE,1:GREEN,3:PURPLE,2:ORANGE}
    bw=int(55*scale); bh=int(52*scale); gap=int(22*scale)
    xx=x
    for i,v in enumerate(vals):
        c=cols[v]
        box(xx,y,xx+bw,y+bh,fill=(248,248,248),outline=c,width=3,r=7)
        center_text(str(v),(xx+bw/2,y+bh/2),fnt(int(28*scale),True))
        if i < len(vals)-1:
            arrow(xx+bw+3,y+bh/2,xx+bw+gap-3,y+bh/2,width=2,head=8)
        xx += bw+gap
    return xx

def draw_brain(x,y,w,h,parcel=False):
    # Hand-drawn schematic, intentionally simple and non-generative.
    pts=[(x+.04*w,y+.58*h),(x+.10*w,y+.33*h),(x+.24*w,y+.17*h),(x+.47*w,y+.11*h),
         (x+.72*w,y+.18*h),(x+.90*w,y+.35*h),(x+.95*w,y+.58*h),(x+.82*w,y+.76*h),
         (x+.60*w,y+.83*h),(x+.35*w,y+.78*h),(x+.16*w,y+.72*h)]
    d.polygon(pts,fill=(205,205,205) if not parcel else (225,225,225),outline=(65,65,65))
    if parcel:
        colors=[RED,GREEN,PURPLE,ORANGE,TEAL,BLUE,(210,80,160),(120,90,50)]
        centers=[(.25,.36),(.45,.27),(.63,.34),(.78,.48),(.55,.55),(.30,.60),(.68,.67),(.43,.70)]
        for (cx,cy),c in zip(centers,colors):
            r=.14*w
            d.ellipse([x+cx*w-r/2,y+cy*h-r/2,x+cx*w+r/2,y+cy*h+r/2],fill=c,outline=(60,60,60),width=1)
    else:
        for cx,cy,rr in [(.42,.42,.12),(.54,.48,.09),(.68,.60,.10),(.34,.60,.08)]:
            d.ellipse([x+(cx-rr/2)*w,y+(cy-rr/2)*h,x+(cx+rr/2)*w,y+(cy+rr/2)*h],fill=(240,75,25),outline=(205,40,20))
        # sulcal curves
        for k in range(5):
            yy=y+(0.28+0.09*k)*h
            d.arc([x+.14*w,yy-.06*h,x+.84*w,yy+.16*h],200,345,fill=(110,110,110),width=2)

def draw_meg_scanner(x,y,w,h):
    # Neutral gray MEG-chair schematic.
    d.rounded_rectangle([x+.22*w,y+.03*h,x+.70*w,y+.53*h],radius=int(.08*w),fill=(232,232,232),outline=(150,150,150),width=2)
    d.pieslice([x+.10*w,y+.12*h,x+.72*w,y+.75*h],70,292,fill=(242,242,242),outline=(165,165,165),width=2)
    d.rectangle([x+.43*w,y+.48*h,x+.66*w,y+.86*h],fill=(195,195,195),outline=(120,120,120),width=2)
    d.polygon([(x+.40*w,y+.76*h),(x+.72*w,y+.76*h),(x+.80*w,y+.98*h),(x+.28*w,y+.98*h)],fill=(130,130,130),outline=(90,90,90))
    d.line([x+.35*w,y+.97*h,x+.18*w,y+h],fill=(80,80,80),width=3)

def draw_eeg_head(x,y,w,h):
    # Head profile with blue 128/256-channel cap dots.
    d.ellipse([x+.20*w,y+.06*h,x+.78*w,y+.82*h],fill=(232,232,232),outline=(170,170,170),width=2)
    d.polygon([(x+.72*w,y+.35*h),(x+.93*w,y+.47*h),(x+.72*w,y+.52*h)],fill=(232,232,232),outline=(170,170,170))
    d.arc([x+.55*w,y+.40*h,x+.72*w,y+.58*h],70,280,fill=(190,190,190),width=2)
    for i in range(9):
        for j in range(7):
            cx=x+.24*w+i*.058*w + (j%2)*.025*w
            cy=y+.11*h+j*.075*h
            if ((cx-x-.49*w)/(.32*w))**2+((cy-y-.38*h)/(.34*h))**2<1.15:
                d.ellipse([cx-4,cy-4,cx+4,cy+4],fill=BLUE)
    d.ellipse([x+.60*w,y+.50*h-5,x+.60*w+10,y+.50*h+5],fill=(230,70,180))

def trace(x,y,w,h,color=BLACK,peaks=7,seed=1):
    import random
    rng=random.Random(seed)
    pts=[]; n=220
    for i in range(n):
        xx=x+w*i/(n-1)
        base=y+h*.70 + rng.uniform(-.015,.015)*h
        val=0.0
        for p in range(peaks):
            loc=(p+1)/(peaks+1)+rng.uniform(-.015,.015)
            amp=.22+rng.uniform(0,.18)
            sig=.009+rng.uniform(0,.008)
            t=i/(n-1)
            val += amp*math.exp(-((t-loc)**2)/(2*sig**2))
        yy=base-val*h
        pts.append((xx,yy))
    d.line(pts,fill=color,width=2)

def tgpac(x,y,w,h):
    pts=[]
    for i in range(300):
        t=i/299
        yy=y+h*.42-math.sin(t*math.pi*2*3.4)*h*.22
        pts.append((x+t*w,yy))
    d.line(pts,fill=(25,110,230),width=4)
    pts=[]
    for i in range(1000):
        t=i/999
        envelope=(.14+.26*(.5+.5*math.sin(t*math.pi*2*3.4)))
        yy=y+h*.55-math.sin(t*math.pi*2*40)*h*envelope
        pts.append((x+t*w,yy))
    d.line(pts,fill=(30,200,70),width=2)

# Layout
margin=18; labelw=170
table_x=margin+labelw
content_w=W-table_x-margin
colw=content_w/4
top=20
header_h=75
obj_h=105; data_h=235; paradigm_h=190; analysis_h=465; output_h=95
aim_y=top
grid_y=aim_y+header_h+20
row_heights=[obj_h,data_h,paradigm_h,analysis_h,output_h]
row_names=["OBJECTIVE","DATA","PARADIGM","ANALYSIS","OUTPUT"]

# section headers
box(table_x,aim_y,table_x+2*colw-8,aim_y+header_h,outline=BLACK,width=3,r=8)
box(table_x+2*colw+8,aim_y,W-margin,aim_y+header_h,outline=BLACK,width=3,r=8)
center_text("AIM 1 | FOUNDATION",(table_x+colw,aim_y+header_h/2),fnt(34,True))
center_text("AIM 2 | TRANSLATION + GENERALIZATION",(table_x+3*colw,aim_y+header_h/2),fnt(32,True))

# grid
yy=grid_y
for rn,rh in zip(row_names,row_heights):
    box(margin,yy,table_x-8,yy+rh,outline=BLACK,width=3)
    center_text(rn,((margin+table_x-8)/2,yy+rh/2),fnt(27,True))
    for c in range(4):
        x0=table_x+c*colw
        box(x0,yy,x0+colw-8,yy+rh,outline=BLACK,width=3)
    yy+=rh

# Aim titles
titles=[("Aim 1A","Benchmark replay inference"),("Aim 1B","Define motor-replay phenotype"),
        ("Aim 2A","Establish EEG-MEG fidelity"),("Aim 2B","Generalize to visual memory")]
for c,(a,b) in enumerate(titles):
    x0=table_x+c*colw; y0=grid_y
    center_text(a,(x0+colw/2-4,y0+34),fnt(34,True))
    center_text(b,(x0+colw/2-4,y0+78),fnt(27,False))

# Data row
data_y=grid_y+obj_h
# A1
draw_brain(table_x+25,data_y+20,colw*.52,data_h-40,False)
text_in("Existing MEG/MRI\nN = 31",table_x+colw*.58,data_y+35,table_x+colw-25,data_y+data_h-30,fnt(26),pad=4)
# A1B
draw_brain(table_x+colw+20,data_y+28,colw*.55,data_h-50,True)
text_in("Same motor cohort\nsource-space MEG",table_x+colw+colw*.60,data_y+35,table_x+2*colw-25,data_y+data_h-30,fnt(25),pad=2)
# A2A
draw_meg_scanner(table_x+2*colw+22,data_y+20,colw*.28,data_h-40)
draw_eeg_head(table_x+2*colw+colw*.32,data_y+18,colw*.30,data_h-40)
center_text("MEG",(table_x+2*colw+colw*.17,data_y+data_h-24),fnt(23,True))
center_text("EEG",(table_x+2*colw+colw*.47,data_y+data_h-24),fnt(23,True))
text_in("Simultaneous\nMEG + 256-ch EEG\nN = 31",table_x+2*colw+colw*.64,data_y+30,table_x+3*colw-20,data_y+data_h-35,fnt(22),pad=2)
# A2B
draw_eeg_head(table_x+3*colw+50,data_y+18,colw*.36,data_h-40)
text_in("128-channel EEG\nN = 31",table_x+3*colw+colw*.48,data_y+35,W-margin-25,data_y+data_h-35,fnt(25),pad=2)

# Paradigm row
par_y=data_y+data_h
seq_boxes(table_x+60,par_y+62,1.05)
center_text("36 ×",(table_x+colw+65,par_y+94),fnt(30,True))
d.line([table_x+colw+145,par_y+45,table_x+colw+145,par_y+145],fill=BLACK,width=3)
for i,t in enumerate(["10 s\npractice","10 s\nrest"]):
    x0=table_x+colw+180+i*175
    box(x0,par_y+43,x0+145,par_y+145,outline=BLACK,width=2,r=8)
    center_text(t,(x0+72,par_y+94),fnt(23))
# same participant
x0=table_x+2*colw+35
box(x0,par_y+62,table_x+3*colw-42,par_y+128,outline=BLACK,width=2,r=8)
center_text("same participant • same task • same rest",((x0+table_x+3*colw-42)/2,par_y+95),fnt(23,True))
# visual sequences
vx=table_x+3*colw+28
center_text("Visual sequence A",(vx+125,par_y+28),fnt(20,True)); center_text("Visual sequence B",(vx+420,par_y+28),fnt(20,True))
shapesA=[("tri",RED),("circ",BLUE),("sq",GREEN),("star",ORANGE)]
shapesB=[("diamond",PURPLE),("plus",GRAY),("heart",(240,40,160)),("hex",ORANGE)]
def icon_box(x,y,kind,color):
    box(x,y,x+58,y+58,outline=BLACK,width=2,r=6)
    cx=x+29; cy=y+29
    if kind=="tri":
        d.polygon([(cx,cy-17),(cx-18,cy+15),(cx+18,cy+15)],fill=color)
    elif kind=="circ": d.ellipse([cx-17,cy-17,cx+17,cy+17],fill=color)
    elif kind=="sq": d.rectangle([cx-16,cy-16,cx+16,cy+16],fill=color)
    elif kind=="diamond": d.polygon([(cx,cy-18),(cx-18,cy),(cx,cy+18),(cx+18,cy)],fill=color)
    elif kind=="plus":
        d.rectangle([cx-7,cy-18,cx+7,cy+18],fill=color); d.rectangle([cx-18,cy-7,cx+18,cy+7],fill=color)
    elif kind=="heart":
        d.ellipse([cx-16,cy-14,cx,cy+2],fill=color);d.ellipse([cx,cy-14,cx+16,cy+2],fill=color);d.polygon([(cx-17,cy-5),(cx+17,cy-5),(cx,cy+19)],fill=color)
    elif kind=="star":
        pts=[]
        for k in range(10):
            r=18 if k%2==0 else 8; a=-math.pi/2+k*math.pi/5
            pts.append((cx+r*math.cos(a),cy+r*math.sin(a)))
        d.polygon(pts,fill=color)
    elif kind=="hex":
        pts=[(cx+18*math.cos(k*math.pi/3),cy+18*math.sin(k*math.pi/3)) for k in range(6)]
        d.polygon(pts,fill=color)
for i,(s,c) in enumerate(shapesA): icon_box(vx+i*66,par_y+48,s,c)
for i,(s,c) in enumerate(shapesB): icon_box(vx+295+i*66,par_y+48,s,c)
center_text("independent item localizer",(vx+265,par_y+120),fnt(19))
for i,t in enumerate(["36 cycles","10 s learn","10 s rest"]):
    x=vx+32+i*185
    box(x,par_y+139,x+155,par_y+178,outline=BLACK,width=2,r=6); center_text(t,(x+77,par_y+158),fnt(18))

# Analysis row
ana_y=par_y+paradigm_h
# A1
x=table_x+20
d.text((x+0,ana_y+15),"Known-sequence injection",font=fnt(27,True),fill=BLACK)
seq_boxes(x+0,ana_y+72,.72)
arrow(x+140,ana_y+138,x+140,ana_y+180,width=3,head=10)
trace(x+0,ana_y+175,260,90,BLACK,peaks=7,seed=2)
d.text((x+285,ana_y+72),"Replay probability",font=fnt(20),fill=BLACK)
trace(x+285,ana_y+98,210,70,BLACK,peaks=7,seed=3)
d.text((x+285,ana_y+175),"Compression",font=fnt(20),fill=BLACK)
trace(x+285,ana_y+202,85,60,BLACK,peaks=6,seed=4); arrow(x+380,ana_y+235,x+410,ana_y+235,width=2,head=8); trace(x+420,ana_y+205,75,55,BLACK,peaks=4,seed=5)
for i,t in enumerate(["Whole-sequence","TDLM"]):
    bx=x+i*285
    box(bx,ana_y+286,bx+270,ana_y+340,outline=BLACK,width=2,r=7); center_text(t,(bx+135,ana_y+313),fnt(21))
center_text("sensitivity + false positives + timing/compression bias",(x+270,ana_y+380),fnt(18))
d.line([x,ana_y+405,x+540,ana_y+405],fill=BLACK,width=2)
center_text("frozen visual-memory MEG robustness test",(x+270,ana_y+435),fnt(18))

# A1B
x=table_x+colw+25
d.text((x,ana_y+15),"Source-to-parcel replay mapping",font=fnt(26,True),fill=BLACK)
draw_brain(x+15,ana_y+60,220,135,False); arrow(x+235,ana_y+128,x+300,ana_y+128,width=3,head=12); draw_brain(x+305,ana_y+60,225,135,True)
seq_boxes(x+10,ana_y+205,.74); arrow(x+365,ana_y+230,x+405,ana_y+230,width=2,head=8)
for i,c in enumerate([BLUE,GREEN,PURPLE,ORANGE,BLUE]): d.rectangle([x+420+i*22,ana_y+208,x+436+i*22,ana_y+254],fill=c,outline=BLACK,width=1)
center_text("Theta-Gamma PAC",(x+270,ana_y+285),fnt(25,True)); tgpac(x+25,ana_y+300,500,105)
center_text("compression + fidelity + F/R",(x+270,ana_y+420),fnt(19))
center_text("HIP ↔ NEO timing",(x+270,ana_y+446),fnt(19))
center_text("micro-offline learning",(x+270,ana_y+472),fnt(19))

# A2A
x=table_x+2*colw+18
d.text((x,ana_y+15),"Cross-modality correspondence",font=fnt(26,True),fill=BLACK)
draw_meg_scanner(x+55,ana_y+65,160,150); draw_eeg_head(x+310,ana_y+55,170,165)
arrow(x+225,ana_y+130,x+300,ana_y+130,width=3,head=12)
center_text("MEG replay",(x+135,ana_y+220),fnt(20)); center_text("EEG replay",(x+395,ana_y+220),fnt(20))
trace(x+5,ana_y+235,260,85,BLACK,peaks=8,seed=6); trace(x+285,ana_y+235,260,85,BLACK,peaks=8,seed=7)
for i,t in enumerate(["strength","specificity","compression","TGPAC"]):
    bw=128; bx=x+i*136
    box(bx,ana_y+326,bx+bw,ana_y+378,outline=BLACK,width=2,r=7); center_text(t,(bx+bw/2,ana_y+352),fnt(18))
d.text((x+0,ana_y+402),"Montage reduction",font=fnt(19),fill=BLACK)
for i,t in enumerate(["256","128","61"]):
    bx=x+285+i*105
    box(bx,ana_y+393,bx+72,ana_y+444,outline=BLACK,width=2,r=6); center_text(t,(bx+36,ana_y+418),fnt(22,True))
    if i<2: arrow(bx+75,ana_y+418,bx+100,ana_y+418,width=2,head=7)
center_text("independent 61-ch motor EEG, N = 64",(x+270,ana_y+470),fnt(18))

# A2B
x=table_x+3*colw+20
d.text((x,ana_y+15),"Visual-memory replay analysis",font=fnt(26,True),fill=BLACK)
for i,(s,c) in enumerate(shapesA): icon_box(x+i*70,ana_y+78,s,c)
arrow(x+290,ana_y+107,x+345,ana_y+107,width=3,head=10)
for i,c in enumerate([RED,BLUE,GREEN,ORANGE]): d.rectangle([x+365+i*38,ana_y+70,x+390+i*38,ana_y+145],fill=c,outline=BLACK,width=1)
labels=["target vs alternate /\nreverse / shuffled","TGPAC","behavior"]; widths=[250,125,125]; bx=x
for i,(t,bw) in enumerate(zip(labels,widths)):
    box(bx,ana_y+205,bx+bw,ana_y+285,outline=BLACK,width=2,r=7)
    center_text(t,(bx+bw/2,ana_y+245),fnt(18))
    bx += bw+15
d.line([x,ana_y+330,x+535,ana_y+330],fill=BLACK,width=2)
center_text("separate 4-min positive-control rest",(x+267,ana_y+365),fnt(18))

# Output row
out_y=ana_y+analysis_h
outputs=["Calibrated operating envelope","Physiological replay phenotype","Scalable EEG replay phenotype","Cross-domain generalization"]
for c,t in enumerate(outputs):
    center_text(t,(table_x+c*colw+(colw-8)/2,out_y+output_h/2),fnt(27,True))

img.save(FIG, quality=95)

# ---------- DOCX reconstruction ----------
raw = SOURCE.read_text(encoding="utf-8", errors="replace").replace("\r\n","\n")
# Remove flattened overall-strategy summary table; the figure replaces it.
raw = re.sub(r"(Overall strategy\..*?serial dependency\.)\s*Component.*?(?=C\.1\. Preliminary foundation and feasibility)", r"\1\n\n", raw, flags=re.S)
# Replace preliminary placeholders with a token to rebuild as a real table.
raw = re.sub(r"PRELIMINARY DATA PLACEHOLDER A.*?Figure placeholder\. Final panels will be assembled only from empirical study outputs; no illustrative or generated imagery will be used\.", "[[PRELIM_TABLE]]", raw, flags=re.S)
# Replace milestone flattened table with a token.
raw = re.sub(r"Project year\s+Scientific milestones\s+Undergraduate research participation.*?(?=Rigor and dissemination\.)", "[[MILESTONE_TABLE]]\n", raw, flags=re.S)

# normalize paragraphs: double newline boundaries; single linebreaks inside block -> space
blocks=[]
for b in re.split(r"\n\s*\n", raw):
    b=b.strip()
    if not b: continue
    if b in ("[[PRELIM_TABLE]]","[[MILESTONE_TABLE]]"):
        blocks.append(b); continue
    b=re.sub(r"\s*\n\s*"," ",b)
    b=re.sub(r"\s+"," ",b).strip()
    blocks.append(b)

doc=Document()
sec=doc.sections[0]
sec.top_margin=Inches(.5); sec.bottom_margin=Inches(.5); sec.left_margin=Inches(.5); sec.right_margin=Inches(.5)
sec.header_distance=Inches(.2); sec.footer_distance=Inches(.2)

styles=doc.styles
normal=styles["Normal"]
normal.font.name="Arial"; normal.font.size=Pt(11)
normal._element.rPr.rFonts.set(qn("w:eastAsia"),"Arial")
normal.paragraph_format.space_after=Pt(0); normal.paragraph_format.space_before=Pt(0)
normal.paragraph_format.line_spacing=1.0

def set_run_font(run,size=11,bold=None,italic=None):
    run.font.name="Arial"; run._element.rPr.rFonts.set(qn("w:eastAsia"),"Arial")
    run.font.size=Pt(size)
    if bold is not None: run.bold=bold
    if italic is not None: run.italic=italic

def set_cell_margins(cell, top=60, start=70, bottom=60, end=70):
    tc=cell._tc; tcPr=tc.get_or_add_tcPr(); tcMar=tcPr.first_child_found_in("w:tcMar")
    if tcMar is None:
        tcMar=OxmlElement("w:tcMar"); tcPr.append(tcMar)
    for m,v in [("top",top),("start",start),("bottom",bottom),("end",end)]:
        node=tcMar.find(qn("w:"+m))
        if node is None: node=OxmlElement("w:"+m); tcMar.append(node)
        node.set(qn("w:w"),str(v)); node.set(qn("w:type"),"dxa")

def shade_cell(cell, fill):
    tcPr=cell._tc.get_or_add_tcPr(); shd=OxmlElement("w:shd"); shd.set(qn("w:fill"),fill); tcPr.append(shd)

def set_cell_text(cell,text,size=8.5,bold=False,align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text=""
    p=cell.paragraphs[0]; p.alignment=align; p.paragraph_format.space_after=Pt(0); p.paragraph_format.line_spacing=1.0
    r=p.add_run(text); set_run_font(r,size,bold)
    cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_margins(cell)

def add_prelim_table():
    t=doc.add_table(rows=2,cols=2); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
    data=[
      ("PRELIMINARY DATA PLACEHOLDER A","Reproduced MEG motor-replay phenotype and replay-behavior relationship.\nInsert empirical panel and exact N/effect size only."),
      ("PRELIMINARY DATA PLACEHOLDER B","Motor action versus sequence-context representation result.\nInsert published empirical panel only."),
      ("PRELIMINARY DATA PLACEHOLDER C","Constricted-pupil replay state and rest TGPAC-learning preliminary result.\nInsert finalized empirical panels only."),
      ("PRELIMINARY DATA PLACEHOLDER D","Reproduced visual-sequence EEG replay analysis and injection/compression calibration.\nInsert finalized empirical panels only.")
    ]
    for i,(h,b) in enumerate(data):
        cell=t.cell(i//2,i%2); cell.width=Inches(3.55); set_cell_margins(cell,80,100,80,100)
        p=cell.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        r=p.add_run(h); set_run_font(r,8.2,True)
        p2=cell.add_paragraph(); p2.alignment=WD_ALIGN_PARAGRAPH.CENTER; p2.paragraph_format.space_after=Pt(0)
        r=p2.add_run(b); set_run_font(r,7.8)
        cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p=doc.add_paragraph()
    p.alignment=WD_ALIGN_PARAGRAPH.LEFT; p.paragraph_format.space_after=Pt(0)
    r=p.add_run("Figure placeholder. Final panels will be assembled only from empirical study outputs; no illustrative or generated imagery will be used.")
    set_run_font(r,8)

def add_milestone_table():
    t=doc.add_table(rows=4,cols=3); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.style="Table Grid"; t.autofit=False
    headers=["Project year","Scientific milestones","Undergraduate research participation"]
    for j,h in enumerate(headers):
        set_cell_text(t.cell(0,j),h,8,True); shade_cell(t.cell(0,j),"E7E7E7")
    rows=[
      ("Year 1","Complete Aim 1A operating envelope; complete primary Aim 1B source/TGPAC analyses; freeze shared replay definitions; initiate Aim 2A acquisition.","Existing-data QC, injection benchmarking, reproducible code review; assist EEG setup and synchronized acquisition."),
      ("Year 2","Complete N=31 Aim 2A acquisition/analysis; complete 256-to-128/61 density reductions and N=64 external 61-channel analysis; initiate Aim 2B.","EEG preprocessing/QC, montage reduction, external-dataset replication, supervised 128-channel acquisition."),
      ("Year 3","Complete N=31 Aim 2B; integrate motor/visual phenotype; finalize manuscripts, public code, and conference dissemination.","Lead supervised Aim 2B data workflow; contribute figures, reproducible reports, abstracts, and manuscript sections.")
    ]
    widths=[.75,3.35,3.35]
    for i,row in enumerate(rows,1):
        for j,val in enumerate(row):
            t.cell(i,j).width=Inches(widths[j]); set_cell_text(t.cell(i,j),val,7.8,False)
    return t

section_heads={"A. SIGNIFICANCE","B. INNOVATION","C. APPROACH"}
sub_heads={"C.1. Preliminary foundation and feasibility","C.2. Common rigor, reproducibility, and statistical framework","C.3. Integration, undergraduate participation, milestones, and dissemination"}
aim_heads=[
 "Aim 1A. Benchmark noninvasive motor-replay inference",
 "Aim 1B. Define the physiological motor-replay signature",
 "Aim 2A. Establish EEG-MEG cross-modal fidelity",
 "Aim 2B. Test generalizability of the EEG replay signature to visual sequence memory",
]

inserted_figure=False
for b in blocks:
    if b=="[[PRELIM_TABLE]]":
        add_prelim_table(); continue
    if b=="[[MILESTONE_TABLE]]":
        add_milestone_table(); continue

    if b in section_heads:
        p=doc.add_paragraph(); p.paragraph_format.space_before=Pt(4); p.paragraph_format.space_after=Pt(2)
        r=p.add_run(b); set_run_font(r,11,True)
        continue
    if b in sub_heads or b in aim_heads:
        p=doc.add_paragraph(); p.paragraph_format.space_before=Pt(3); p.paragraph_format.space_after=Pt(0)
        r=p.add_run(b); set_run_font(r,11,True)
        continue

    p=doc.add_paragraph()
    p.paragraph_format.space_after=Pt(0); p.paragraph_format.line_spacing=1.0
    # Bold the lead phrase up to the first period for labeled paragraphs.
    labels=[
      "Significance of the problem.","Scientific premise.","Validity gap.","Representational-content gap.",
      "Physiological-organization gap.","Scalability and generalizability gap.","Expected contribution and impact.",
      "AREA impact.","Conceptual innovation: replay as a calibrated physiological phenotype.",
      "Validation innovation: known ground truth in realistic human resting-brain activity.",
      "Representational and physiological innovation: defining what is replayed and when it occurs.",
      "Cross-modal innovation: translating a biological phenotype rather than a decoder.",
      "Cross-domain innovation: matched brief-rest replay in visual sequence memory.","Overall innovative advance.",
      "Overall strategy.","Analysis separation and avoidance of circularity.","Complementary replay measures.",
      "Controls and realistic nulls.","Primary endpoints, multiplicity, and estimation.","Behavior and reproducibility.",
      "Rationale and hypothesis.","Dataset and representational hierarchy.","Replay inference and controls.",
      "Known-ground-truth injection.","Cross-domain robustness and interpretation.","Potential problems and alternatives.",
      "Source-space replay phenotype.","TGPAC and behavioral relevance.","Expected outcomes and alternatives.",
      "Participants and protocol.","Simultaneous acquisition and independent preprocessing.",
      "Cross-modal fidelity and density robustness.","Power and interpretation.","Participants and adapted visual-sequence paradigm.",
      "EEG replay, physiology, and behavior.","Power and alternatives.","Aim integration and independence.",
      "Undergraduate participation and supervision.","Milestones and decision rules.","Rigor and dissemination."
    ]
    matched=None
    for lab in labels:
        if b.startswith(lab):
            matched=lab; break
    if matched:
        r=p.add_run(matched+" "); set_run_font(r,11,True)
        r=p.add_run(b[len(matched):].lstrip()); set_run_font(r,11)
    else:
        r=p.add_run(b); set_run_font(r,11)

    if b.startswith("Overall strategy.") and not inserted_figure:
        pfig=doc.add_paragraph(); pfig.alignment=WD_ALIGN_PARAGRAPH.CENTER; pfig.paragraph_format.space_before=Pt(3); pfig.paragraph_format.space_after=Pt(1)
        pfig.add_run().add_picture(str(FIG), width=Inches(7.25))
        cap=doc.add_paragraph(); cap.alignment=WD_ALIGN_PARAGRAPH.LEFT; cap.paragraph_format.space_after=Pt(2)
        r=cap.add_run("Figure 1. Overall strategy. "); set_run_font(r,8,True)
        r=cap.add_run("Aim 1 calibrates replay inference and defines the physiological motor-replay phenotype; Aim 2 tests same-rest MEG-EEG fidelity and cross-domain generalizability to visual sequence memory. Each subaim yields an independently interpretable result."); set_run_font(r,8)
        inserted_figure=True

# Ensure all styles default to Arial.
for st in styles:
    try:
        st.font.name="Arial"; st._element.rPr.rFonts.set(qn("w:eastAsia"),"Arial")
    except Exception:
        pass

doc.save(DOCX)
print(DOCX)
